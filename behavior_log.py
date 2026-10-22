"""Ghi hành vi từng học sinh theo thời gian và xuất file Excel.

Lưu gọn để chạy video dài nhiều giờ:
  - đợt hành vi: (track, hành vi) liên tiếp được gộp thành 1 dòng bắt đầu → kết thúc
  - thống kê theo track: số lần quan sát, điểm, điểm theo phút
Thời gian lưu bằng chỉ số khung, đổi ra ngày giờ lúc xuất qua VideoClock
(mốc OSD có thể đọc được sau khi đã ghi).
"""

from __future__ import annotations

import io
import re
import threading
from collections import Counter, defaultdict
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import timedelta
from pathlib import Path

from behavior import POSITIVE_BEHAVIORS as POSITIVE
from behavior import REASON_SHORT
from video_clock import VideoClock

BEHAVIORS = tuple(name for name, _short in REASON_SHORT)
SHORT_OF = dict(REASON_SHORT)
# Không thấy lại hành vi sau N giây video → kết thúc đợt
EPISODE_GAP_SEC = 3.0
# Số lần phân tích nhận ra cùng 1 mã HS mới tính là có mặt (tránh 1 lần khớp nhầm)
PRESENT_MIN_FRAMES = 3
DASHBOARD_MAX_MINUTES = 90
# Mức tin cậy điểm danh (theo % khớp ArcFace cao nhất và số lần nhận ra)
CONF_HIGH_PCT = 45.0
CONF_HIGH_VOTES = 10
CONF_MID_PCT = 35.0
FACE_THUMB_H = 96  # px, ảnh mặt nhúng vào Excel
FACE_CROP_FRAC = 0.45  # phần trên box người chứa mặt

_SECONDS_RE = re.compile(r"(\d+(?:\.\d+)?)s\)")


def _normalize(reason: str) -> tuple[str | None, float]:
    """'Cúi đầu lâu (3.2s)' → ('Cúi đầu lâu', 3.2)."""
    for name in BEHAVIORS:
        if reason.startswith(name):
            m = _SECONDS_RE.search(reason)
            return name, float(m.group(1)) if m else 0.0
    return None, 0.0


@dataclass
class _Episode:
    tid: int
    behavior: str
    start_frame: int
    last_frame: int
    start_wall: float
    last_wall: float
    min_score: int
    detail: str = ""


@dataclass
class _TrackStats:
    n_obs: int = 0
    sum_score: int = 0
    min_score: int = 100
    unfocused_obs: int = 0
    first_frame: int = 0
    last_frame: int = 0
    first_wall: float = 0.0
    last_wall: float = 0.0
    minute_sum: dict[int, int] = field(default_factory=lambda: defaultdict(int))
    minute_n: dict[int, int] = field(default_factory=lambda: defaultdict(int))


class BehaviorRecorder:
    def __init__(self, fps: float, clock: VideoClock, source_label: str,
                 student_names: dict[str, str] | None = None,
                 legend_fn: Callable[[], list[dict]] | None = None,
                 student_classes: dict[str, str] | None = None,
                 default_class: str = ""):
        self.fps = fps
        self.clock = clock
        self.source_label = source_label
        self.legend_fn = legend_fn or (lambda: [])
        self.report_name = "bao_cao_hanh_vi.xlsx"
        self.student_names = student_names or {}
        self.student_classes = student_classes or {}
        self.default_class = default_class
        self._lock = threading.Lock()
        # sid → [tổng %, số lần, % cao nhất] độ khớp ArcFace
        self._sid_pct: dict[str, list[float]] = {}
        # sid → (% khớp, JPEG mặt) ảnh rõ nhất trong lớp
        self._face_best: dict[str, tuple[float, bytes]] = {}
        self._open: dict[tuple[int, str], _Episode] = {}
        self._closed: list[_Episode] = []
        self._stats: dict[int, _TrackStats] = {}
        self._sid_votes: dict[int, Counter] = defaultdict(Counter)
        self._names: dict[str, str] = {}
        self.first_frame: int | None = None
        self.last_frame = 0
        self.first_wall = 0.0
        self.last_wall = 0.0

    @staticmethod
    def _face_jpeg(frame, d: dict) -> bytes | None:
        import cv2

        x1, y1, x2, y2 = (int(d[k]) for k in ("x1", "y1", "x2", "y2"))
        crop = frame[max(0, y1):max(0, y1) + int((y2 - y1) * FACE_CROP_FRAC), max(0, x1):max(0, x2)]
        if crop.size == 0:
            return None
        h, w = crop.shape[:2]
        crop = cv2.resize(crop, (max(1, int(w * FACE_THUMB_H / h)), FACE_THUMB_H),
                          interpolation=cv2.INTER_AREA if h > FACE_THUMB_H else cv2.INTER_CUBIC)
        ok, buf = cv2.imencode(".jpg", crop, [cv2.IMWRITE_JPEG_QUALITY, 90])
        return buf.tobytes() if ok else None

    def record(self, frame_idx: int, detections: list[dict], wall: float, frame=None) -> None:
        gap_frames = EPISODE_GAP_SEC * self.fps
        with self._lock:
            if self.first_frame is None:
                self.first_frame, self.first_wall = frame_idx, wall
            self.last_frame, self.last_wall = frame_idx, wall
            minute = int((frame_idx - self.first_frame) / self.fps // 60)
            seen: set[tuple[int, str]] = set()
            for d in detections:
                tid = int(d["track_id"])
                if d.get("known") and d.get("student_id"):
                    sid = d["student_id"]
                    self._sid_votes[tid][sid] += 1
                    if d.get("name"):
                        self._names[sid] = d["name"]
                    pct = d.get("match_pct")
                    if pct is not None:
                        acc = self._sid_pct.setdefault(sid, [0.0, 0, 0.0])
                        acc[0] += pct
                        acc[1] += 1
                        acc[2] = max(acc[2], pct)
                        if frame is not None and pct > self._face_best.get(sid, (-1.0, b""))[0]:
                            jpg = self._face_jpeg(frame, d)
                            if jpg:
                                self._face_best[sid] = (pct, jpg)
                score = d.get("score")
                if score is None:
                    continue
                st = self._stats.get(tid)
                if st is None:
                    st = self._stats[tid] = _TrackStats(
                        first_frame=frame_idx, first_wall=wall)
                st.n_obs += 1
                st.sum_score += score
                st.min_score = min(st.min_score, score)
                st.unfocused_obs += 0 if d.get("is_focused") else 1
                st.last_frame, st.last_wall = frame_idx, wall
                st.minute_sum[minute] += score
                st.minute_n[minute] += 1

                for reason in d.get("reasons") or []:
                    name, secs = _normalize(reason)
                    if name is None:
                        continue
                    key = (tid, name)
                    seen.add(key)
                    ep = self._open.get(key)
                    if ep is None:
                        back = int(round(secs * self.fps))
                        self._open[key] = _Episode(
                            tid, name, max(0, frame_idx - back), frame_idx,
                            wall - secs, wall, score, reason)
                    else:
                        ep.last_frame, ep.last_wall = frame_idx, wall
                        ep.min_score = min(ep.min_score, score)
                        ep.detail = reason
            for key in list(self._open):
                ep = self._open[key]
                if key not in seen and frame_idx - ep.last_frame > gap_frames:
                    self._closed.append(self._open.pop(key))

    # ---------------------------------------------------------------- xuất

    def _student_of(self, tid: int) -> str | None:
        votes = self._sid_votes.get(tid)
        return votes.most_common(1)[0][0] if votes else None

    def _who(self, tid: int) -> tuple[str, str]:
        sid = self._student_of(tid)
        if sid is None:
            return f"ID{tid} (chưa nhận diện)", ""
        name = self._names.get(sid) or self.student_names.get(sid) or ""
        return sid, "" if name == sid else name

    def _time_cells(self, frame_idx: int, wall: float) -> tuple[str, str, str]:
        """(ngày như video, giờ, vị trí trong video)."""
        dt = self.clock.at(frame_idx, wall)
        offset = str(timedelta(seconds=int(max(0, frame_idx) / self.fps)))
        if dt is None:
            return "", offset, offset
        return self.clock.format_date(dt), dt.strftime("%H:%M:%S"), offset

    def _per_student(self, episodes: list[_Episode]) -> dict[str, dict]:
        """Gộp các track cùng 1 HS. Gọi khi đang giữ self._lock."""
        per_student: dict[str, dict] = {}
        for tid, st in self._stats.items():
            sid, name = self._who(tid)
            agg = per_student.setdefault(sid, {
                "name": name, "n": 0, "sum": 0, "min": 100, "unf": 0,
                "first": (st.first_frame, st.first_wall), "last": (st.last_frame, st.last_wall),
                "count": Counter(), "secs": Counter(),
            })
            agg["n"] += st.n_obs
            agg["sum"] += st.sum_score
            agg["min"] = min(agg["min"], st.min_score)
            agg["unf"] += st.unfocused_obs
            if st.first_frame < agg["first"][0]:
                agg["first"] = (st.first_frame, st.first_wall)
            if st.last_frame > agg["last"][0]:
                agg["last"] = (st.last_frame, st.last_wall)
        for ep in episodes:
            sid, _name = self._who(ep.tid)
            agg = per_student.get(sid)
            if agg is None:
                continue
            agg["count"][ep.behavior] += 1
            agg["secs"][ep.behavior] += max(0.0, (ep.last_frame - ep.start_frame) / self.fps)
        return per_student

    def _sid_vote_totals(self) -> Counter:
        """Số lần nhận ra mỗi mã HS (cộng các track gán cho HS đó). Gọi khi đang giữ self._lock."""
        votes: Counter = Counter()
        for tid, c in self._sid_votes.items():
            sid = self._student_of(tid)
            votes[sid] += c[sid]
        return votes

    def summary(self) -> dict:
        """Số liệu cho Dashboard: sĩ số theo gallery + mức chú ý cả phiên."""
        with self._lock:
            episodes = self._closed + list(self._open.values())
            per_student = self._per_student(episodes)
            votes = self._sid_vote_totals()
            unknown_tracks = sum(1 for tid in self._stats if self._student_of(tid) is None)
            minute_sum: dict[int, int] = defaultdict(int)
            minute_n: dict[int, int] = defaultdict(int)
            for st in self._stats.values():
                for m, s in st.minute_sum.items():
                    minute_sum[m] += s
                    minute_n[m] += st.minute_n[m]
            names = {**self.student_names, **self._names}
            first_frame, first_wall = self.first_frame or 0, self.first_wall

        present_ids = {sid for sid, n in votes.items() if n >= PRESENT_MIN_FRAMES}
        roster = list(self.student_names)
        roster += sorted(present_ids - set(roster))

        students = []
        for sid in sorted(present_ids):
            agg = per_student.get(sid)
            n = agg["n"] if agg else 0
            negatives = sorted(
                ((b, s) for b, s in agg["secs"].items() if b not in POSITIVE) if agg else (),
                key=lambda x: -x[1],
            )[:2]
            students.append({
                "student_id": sid,
                "name": names.get(sid, ""),
                "avg_score": round(agg["sum"] / n, 1) if n else None,
                "min_score": agg["min"] if n else None,
                "focus_pct": round(100.0 * (n - agg["unf"]) / n, 1) if n else None,
                "first_seen": self._time_cells(*agg["first"])[1] if agg else "",
                "top_negative": [
                    {"name": b, "short": SHORT_OF.get(b, b), "count": agg["count"][b], "secs": round(s, 1)}
                    for b, s in negatives
                ],
            })

        total_n = sum(a["n"] for a in per_student.values())
        total_sum = sum(a["sum"] for a in per_student.values())
        total_unf = sum(a["unf"] for a in per_student.values())
        timeline = []
        for m in sorted(minute_n)[-DASHBOARD_MAX_MINUTES:]:
            f = first_frame + int(m * 60 * self.fps)
            timeline.append({
                "label": self._time_cells(f, first_wall + m * 60)[1],
                "avg": round(minute_sum[m] / minute_n[m], 1),
            })

        return {
            "roster_size": len(roster),
            "present": students,
            "absent": [
                {"student_id": sid, "name": names.get(sid, "")}
                for sid in roster if sid not in present_ids
            ],
            "unknown_tracks": unknown_tracks,
            "class_avg": round(total_sum / total_n, 1) if total_n else None,
            "class_focus_pct": round(100.0 * (total_n - total_unf) / total_n, 1) if total_n else None,
            "timeline": timeline,
        }

    def export_excel(self, path: Path) -> Path:
        with self._lock:
            wb = self._build_workbook()
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return path

    @staticmethod
    def _confidence_level(votes: int, max_pct: float | None) -> str:
        if max_pct is None:
            return ""
        if max_pct >= CONF_HIGH_PCT and votes >= CONF_HIGH_VOTES:
            return "Cao"
        if max_pct >= CONF_MID_PCT:
            return "Trung bình"
        return "Thấp — cần kiểm tra"

    def _write_attendance_sheet(self, ws, per_student: dict[str, dict]) -> None:
        """Sheet chính: mỗi học sinh trong danh sách lớp là 1 hàng. Gọi khi đang giữ self._lock."""
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        ws.title = "Điểm danh theo HS"
        votes = self._sid_vote_totals()
        present_ids = {sid for sid, n in votes.items() if n >= PRESENT_MIN_FRAMES}
        roster = list(self.student_names)
        roster += sorted(present_ids - set(roster))
        names = {**self.student_names, **self._names}

        headers = [
            ("STT", 5), ("Mã HS", 10), ("Họ và tên", 22), ("Lớp", 8), ("Trạng thái", 11),
            ("Ngày (như video)", 12), ("Giờ điểm danh (lần đầu nhận diện)", 13),
            ("Lần cuối thấy", 11), ("Thời gian có mặt (phút)", 11),
            ("Số lần nhận diện", 11), ("Độ khớp TB (%)", 10), ("Độ khớp cao nhất (%)", 10),
            ("Mức tin cậy", 13), ("Ảnh khuôn mặt (trong lớp)", 13),
            ("Điểm chú ý TB", 9), ("Điểm chú ý thấp nhất", 9),
            ("% thời gian tập trung", 11), ("% thời gian mất tập trung", 11),
            ("Hành vi mất tập trung nhiều nhất", 18),
        ]
        for b in BEHAVIORS:
            headers += [(f"{b} (lần)", 9), (f"{b} (giây)", 9)]
        headers.append(("Ghi chú (giáo viên)", 24))
        ncol = len(headers)
        col_img = [h for h, _w in headers].index("Ảnh khuôn mặt (trong lớp)") + 1
        last_col = get_column_letter(ncol)

        rows, images = [], []
        n_present = 0
        for i, sid in enumerate(roster, start=1):
            agg = per_student.get(sid)
            present = sid in present_ids
            n_present += present
            name = names.get(sid, "")
            if name == sid:
                name = ""
            row = [i, sid, name, self.student_classes.get(sid) or self.default_class,
                   "Có mặt" if present else "Vắng"]
            if present and agg:
                date_s, from_s, _o = self._time_cells(*agg["first"])
                _d, to_s, _o2 = self._time_cells(*agg["last"])
                minutes = round(max(0, agg["last"][0] - agg["first"][0]) / self.fps / 60, 1)
            else:
                date_s = from_s = to_s = ""
                minutes = None
            pct = self._sid_pct.get(sid)
            avg_pct = round(pct[0] / pct[1], 1) if present and pct and pct[1] else None
            max_pct = round(pct[2], 1) if present and pct else None
            row += [date_s, from_s, to_s, minutes, votes.get(sid, 0) if present else 0,
                    avg_pct, max_pct, self._confidence_level(votes.get(sid, 0), max_pct) if present else "",
                    ""]
            n = agg["n"] if present and agg else 0
            if n:
                negatives = sorted(((b, s) for b, s in agg["secs"].items() if b not in POSITIVE and s > 0),
                                   key=lambda x: -x[1])
                top = f"{negatives[0][0]} ({negatives[0][1]:.0f}s)" if negatives else "—"
                row += [round(agg["sum"] / n, 1), agg["min"], round(100.0 * (n - agg["unf"]) / n, 1),
                        round(100.0 * agg["unf"] / n, 1), top]
                for b in BEHAVIORS:
                    row += [agg["count"][b], round(agg["secs"][b], 1)]
            else:
                row += [None] * (5 + 2 * len(BEHAVIORS))
            row.append("")
            rows.append(row)
            face = self._face_best.get(sid) if present else None
            images.append(face[1] if face else None)

        # --- phần đầu: thông tin buổi học
        title_font = Font(bold=True, size=14, color="1F4E78")
        label_font = Font(bold=True)
        ws.merge_cells(f"A1:{last_col}1")
        ws["A1"] = "BÁO CÁO ĐIỂM DANH VÀ MỨC ĐỘ TẬP TRUNG"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        d1, t1, o1 = self._time_cells(self.first_frame or 0, self.first_wall)
        _d2, t2, o2 = self._time_cells(self.last_frame, self.last_wall)
        all_n = sum(a["n"] for a in per_student.values())
        all_unf = sum(a["unf"] for a in per_student.values())
        all_sum = sum(a["sum"] for a in per_student.values())
        classes = sorted({r[3] for r in rows if r[3]})
        total = len(roster)
        info = [
            ("Lớp", ", ".join(classes)),
            ("Nguồn video", self.source_label),
            ("Ngày (như video)", d1),
            ("Thời gian phân tích", f"{t1} → {t2}  (video {o1} → {o2})"),
            ("Sĩ số", total),
            ("Có mặt", n_present),
            ("Vắng / chưa nhận diện", total - n_present),
            ("Tỉ lệ chuyên cần", f"{100.0 * n_present / total:.1f}%" if total else ""),
            ("Điểm chú ý TB cả lớp", round(all_sum / all_n, 1) if all_n else ""),
            ("% thời gian tập trung cả lớp", f"{100.0 * (all_n - all_unf) / all_n:.1f}%" if all_n else ""),
            ("Track chưa nhận diện được mặt", sum(1 for s in per_student if s.startswith("ID"))),
            ("Quy tắc", f"Có mặt = ArcFace nhận ra mã HS ≥ {PRESENT_MIN_FRAMES} lần. Mức tin cậy Cao: "
                        f"khớp ≥ {CONF_HIGH_PCT:.0f}% và ≥ {CONF_HIGH_VOTES} lần; Trung bình: khớp ≥ "
                        f"{CONF_MID_PCT:.0f}%; còn lại cần giáo viên kiểm tra."),
        ]
        r0 = 3
        for k, (label, value) in enumerate(info):
            r = r0 + k
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=min(ncol, 13))
            ws.cell(row=r, column=1, value=label).font = label_font
            ws.cell(row=r, column=4, value=value).alignment = Alignment(horizontal="left")

        # --- bảng học sinh
        hr = r0 + len(info) + 1
        thin = Side(style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for c, (h, w) in enumerate(headers, start=1):
            cell = ws.cell(row=hr, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.row_dimensions[hr].height = 48

        fills = {
            "Có mặt": PatternFill("solid", fgColor="E2F5E6"),
            "Vắng": PatternFill("solid", fgColor="FDE2E1"),
            "Cao": PatternFill("solid", fgColor="E2F5E6"),
            "Trung bình": PatternFill("solid", fgColor="FFF2CC"),
            "Thấp — cần kiểm tra": PatternFill("solid", fgColor="FCE4D6"),
        }
        img_max_w = FACE_THUMB_H * 3 // 2
        has_images = any(images)
        widest = 0
        for k, (row, jpg) in enumerate(zip(rows, images)):
            r = hr + 1 + k
            for c, v in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = border
                cell.alignment = Alignment(horizontal="left" if c in (2, 3, ncol) else "center",
                                           vertical="center", wrap_text=c in (3, 13, 19, ncol))
            ws.cell(row=r, column=5).fill = fills[row[4]]
            if row[12] in fills:
                ws.cell(row=r, column=13).fill = fills[row[12]]
            if has_images:
                ws.row_dimensions[r].height = FACE_THUMB_H * 0.75 + 4
            if jpg:
                img = XLImage(io.BytesIO(jpg))
                s = min(FACE_THUMB_H / img.height, img_max_w / img.width)
                img.width, img.height = int(img.width * s), int(img.height * s)
                widest = max(widest, img.width)
                ws.add_image(img, f"{get_column_letter(col_img)}{r}")
        if has_images:
            ws.column_dimensions[get_column_letter(col_img)].width = max(13, widest / 7 + 1)
        ws.freeze_panes = ws.cell(row=hr + 1, column=4)
        if rows:
            ws.auto_filter.ref = f"A{hr}:{last_col}{hr + len(rows)}"
        ws.sheet_view.zoomScale = 90
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = f"{hr}:{hr}"

    def _build_workbook(self):
        """Gọi khi đang giữ self._lock (luồng phân tích vẫn ghi song song)."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        episodes = sorted(
            self._closed + list(self._open.values()), key=lambda e: e.start_frame
        )
        stats = self._stats
        first_frame = self.first_frame or 0
        last_frame, first_wall, last_wall = self.last_frame, self.first_wall, self.last_wall

        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E78")
        bad_fill = PatternFill("solid", fgColor="FDE2E1")
        good_fill = PatternFill("solid", fgColor="E2F5E6")

        def write_sheet(ws, headers, rows, widths):
            ws.append(headers)
            for c in ws[1]:
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for r in rows:
                ws.append(r)
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = ws.dimensions

        per_student = self._per_student(episodes)
        self._write_attendance_sheet(wb.active, per_student)

        # --- Chi tiết hành vi
        ws = wb.create_sheet("Chi tiết hành vi")
        rows = []
        for i, ep in enumerate(episodes, start=1):
            sid, name = self._who(ep.tid)
            date_s, start_s, off_s = self._time_cells(ep.start_frame, ep.start_wall)
            _d, end_s, _o = self._time_cells(ep.last_frame, ep.last_wall)
            dur = round(max(0.0, (ep.last_frame - ep.start_frame) / self.fps), 1)
            rows.append([
                i, sid, name, ep.behavior, SHORT_OF.get(ep.behavior, ""),
                "Tích cực" if ep.behavior in POSITIVE else "Mất tập trung",
                date_s, start_s, end_s, dur, ep.min_score, off_s, ep.tid,
            ])
        write_sheet(
            ws,
            ["STT", "Mã HS", "Họ tên", "Hành vi", "Viết tắt (trên video)", "Loại",
             "Ngày (như video)", "Bắt đầu", "Kết thúc", "Thời lượng (giây)",
             "Điểm thấp nhất", "Vị trí trong video", "Track ID"],
            rows,
            [6, 22, 22, 20, 12, 15, 16, 11, 11, 12, 11, 14, 9],
        )
        for row in ws.iter_rows(min_row=2):
            row[5].fill = good_fill if row[5].value == "Tích cực" else bad_fill

        # --- Track chưa nhận diện được mặt (không thuộc danh sách lớp)
        ws2 = wb.create_sheet("Track chưa nhận diện")
        headers = ["Track", "Ngày (như video)", "Xuất hiện từ", "Đến",
                   "Điểm TB", "Điểm thấp nhất", "% thời gian mất tập trung"]
        for b in BEHAVIORS:
            headers += [f"{b} (lần)", f"{b} (giây)"]
        rows2 = []
        unknown = [s for s in per_student if s.startswith("ID")]
        for sid in sorted(unknown, key=lambda s: int(re.sub(r"\D", "", s) or 0)):
            agg = per_student[sid]
            date_s, from_s, _o = self._time_cells(*agg["first"])
            _d, to_s, _o2 = self._time_cells(*agg["last"])
            row = [
                sid.split()[0], date_s, from_s, to_s,
                round(agg["sum"] / agg["n"], 1) if agg["n"] else None,
                agg["min"],
                round(100.0 * agg["unf"] / agg["n"], 1) if agg["n"] else None,
            ]
            for b in BEHAVIORS:
                row += [agg["count"][b], round(agg["secs"][b], 1)]
            rows2.append(row)
        write_sheet(ws2, headers, rows2, [10, 16, 12, 12, 9, 11, 14] + [11] * (2 * len(BEHAVIORS)))

        # --- Điểm theo phút
        ws3 = wb.create_sheet("Điểm theo phút")
        minute_scores: dict[str, dict[int, list[int]]] = defaultdict(lambda: defaultdict(lambda: [0, 0]))
        all_minutes: set[int] = set()
        for tid, st in stats.items():
            sid, _name = self._who(tid)
            for m, s in st.minute_sum.items():
                cell = minute_scores[sid][m]
                cell[0] += s
                cell[1] += st.minute_n[m]
                all_minutes.add(m)
        minutes = sorted(all_minutes)
        headers3 = ["Mã HS"]
        for m in minutes:
            f = first_frame + int(m * 60 * self.fps)
            _d, t_s, _o = self._time_cells(f, first_wall + m * 60)
            headers3.append(t_s)
        rows3 = []
        for sid in sorted(minute_scores, key=lambda s: (s.startswith("ID"), s)):
            row = [sid]
            for m in minutes:
                tot, n = minute_scores[sid].get(m, (0, 0))
                row.append(round(tot / n) if n else None)
            rows3.append(row)
        write_sheet(ws3, headers3, rows3, [22] + [10] * len(minutes))

        # --- Thông tin
        ws4 = wb.create_sheet("Thông tin")
        src = {
            "osd": "Đọc từ ngày giờ in trên video (OSD)",
            "system": "Đồng hồ máy (camera trực tiếp)",
            "video": "Không đọc được ngày giờ trên video — dùng vị trí trong video",
        }[self.clock.source]
        d1, t1, o1 = self._time_cells(first_frame, first_wall)
        d2, t2, o2 = self._time_cells(last_frame, last_wall)
        info = [
            ("Nguồn video", self.source_label),
            ("Nguồn thời gian", src),
            ("Bắt đầu ghi", f"{d1} {t1}".strip()),
            ("Kết thúc ghi", f"{d2} {t2}".strip()),
            ("Đoạn video đã phân tích", f"{o1} → {o2}"),
            ("Số đợt hành vi", len(episodes)),
            ("Số học sinh / track", len(per_student)),
            ("Ghi chú", "Điểm tập trung 0–100; dưới 50 = mất tập trung. "
                        "Một đợt kết thúc khi không thấy lại hành vi sau "
                        f"{EPISODE_GAP_SEC:.0f} giây video."),
        ]
        for k, v in info:
            ws4.append([k, v])
        for c in ws4["A"]:
            c.font = Font(bold=True)
        ws4.column_dimensions["A"].width = 26
        ws4.column_dimensions["B"].width = 90

        # --- Chú thích nhãn viết tắt
        ws5 = wb.create_sheet("Chú thích")
        rows5 = [
            [lg["short"], lg["name"], "Tích cực" if lg["positive"] else "Mất tập trung",
             lg["condition"], -lg["penalty"] if lg["penalty"] else 0,
             "Bật" if lg.get("enabled", True) else "Tắt"]
            for lg in self.legend_fn()
        ]
        write_sheet(
            ws5,
            ["Viết tắt (trên video)", "Hành vi đầy đủ", "Loại", "Điều kiện nhận biết", "Điểm trừ",
             "Trạng thái (lúc xuất)"],
            rows5,
            [16, 22, 15, 70, 10, 14],
        )
        off_fill = PatternFill("solid", fgColor="E7E6E6")
        for row in ws5.iter_rows(min_row=2):
            row[2].fill = good_fill if row[2].value == "Tích cực" else bad_fill
            if row[5].value == "Tắt":
                row[5].fill = off_fill
        ws5.append([])
        box_notes = [
            ("Khung xanh lá", "Tập trung (điểm ≥ 50)", "E2F5E6"),
            ("Khung đỏ", "Mất tập trung (điểm < 50); đỏ dày = cúi đầu lâu", "FDE2E1"),
            ("Khung xanh dương nhạt", "Chưa lấy được tư thế — chưa chấm điểm", "DDEBF7"),
            ("HS012 85 · Cúi đầu lâu", "Mã học sinh · điểm tập trung (0–100) · hành vi chính", None),
            ("ID7", "Chưa nhận diện được học sinh (số theo dõi tạm)", None),
        ]
        for label, meaning, color in box_notes:
            ws5.append([label, meaning])
            if color:
                ws5.cell(row=ws5.max_row, column=1).fill = PatternFill("solid", fgColor=color)
            ws5.cell(row=ws5.max_row, column=1).font = Font(bold=True)
        return wb
